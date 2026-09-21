"""
Multivariate Exploratory Data Analysis.

Class-based engine covering:
- Multicollinearity / VIF, PCA, decision trees
- Linear / Logistic models, Random Forest
- Two-way / three-way interactions
- K-Means clustering, Isolation Forest & Mahalanobis outliers

All statistical calculations are deterministic. LLM summarization is separate.
"""

from __future__ import annotations

import itertools
import json
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from scipy.stats import chi2
from sklearn.cluster import KMeans
from sklearn.compose import ColumnTransformer
from sklearn.decomposition import PCA
from sklearn.ensemble import IsolationForest, RandomForestClassifier, RandomForestRegressor
from sklearn.inspection import permutation_importance
from sklearn.linear_model import LinearRegression, LogisticRegression
from sklearn.metrics import (
    accuracy_score, mean_absolute_error, mean_squared_error, r2_score,
    silhouette_score, calinski_harabasz_score, davies_bouldin_score,
)
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.tree import DecisionTreeClassifier, DecisionTreeRegressor

from src.config import BASE_DIR
from src.Utils.exception_handler import CustomException
from src.Utils.logger_setup import get_log, track_performance


class MultivariateEDA:
    """Comprehensive multivariate EDA engine; results stored as Excel worksheets."""

    WORKBOOK_NAME = "multivariate_eda.xlsx"
    CLASSIFIED_COLUMNS_PATH = BASE_DIR / "src/Data/classified_columns.json"
    DEFAULT_OUTPUT_DIR = (
        BASE_DIR / "src/EDA_Summary/Multivariate_Analysis/Reports/excel_reports"
    )
    DEFAULT_LOG_DIR = BASE_DIR / "src/EDA_Summary/Multivariate_Analysis/logs"
    ID_NAME_TOKENS = {
        "id", "identifier", "uuid", "customer_id", "user_id", "account_id",
        "transaction_id", "order_id", "product_id", "session_id",
    }
    MAHALANOBIS_CONFIDENCE = 0.975

    def __init__(
        self,
        df: pd.DataFrame | None = None,
        output_dir: Path | None = None,
        target: str | None = None,
        test_size: float = 0.20,
        random_state: int = 42,
        max_depth: int = 5,
        min_samples_leaf: int = 20,
        n_estimators: int = 200,
        permutation_repeats: int = 5,
        vif_threshold: float = 10,
        max_pca_components: int | None = None,
        max_interaction_features: int = 10,
        max_three_way_combinations: int = 100,
        n_clusters: int = 5,
        max_clustering_features: int = 20,
        contamination: float = 0.01,
        summary_top_n: int = 15,
        min_group_size: int = 10,
    ) -> None:
        self.df = df.copy() if df is not None else None
        self.output_dir = Path(output_dir) if output_dir is not None else self.DEFAULT_OUTPUT_DIR
        self.logs_dir = self.DEFAULT_LOG_DIR

        self.target = target
        self.test_size = test_size
        self.random_state = random_state
        self.max_depth = max_depth
        self.min_samples_leaf = min_samples_leaf
        self.n_estimators = n_estimators
        self.permutation_repeats = permutation_repeats
        self.vif_threshold = vif_threshold
        self.max_pca_components = max_pca_components
        self.max_interaction_features = max_interaction_features
        self.max_three_way_combinations = max_three_way_combinations
        self.n_clusters = n_clusters
        self.max_clustering_features = max_clustering_features
        self.contamination = contamination
        self.summary_top_n = summary_top_n
        self.min_group_size = min_group_size

        self.logger = get_log("MultivariateEDA", log_dir=self.logs_dir)
        self.classified_columns: dict[str, Any] = {}
        self.numeric_features: list[str] = []
        self.categorical_features: list[str] = []
        self.top_features: list[str] = []

        self._validate_configuration()
        if self.df is not None:
            self._initialize_features()

    # ------------------------------------------------------------------
    # Configuration
    # ------------------------------------------------------------------

    def _validate_configuration(self) -> None:
        if not 0 < self.test_size < 1:
            raise ValueError("test_size must be between 0 and 1.")
        if self.max_depth < 1:
            raise ValueError("max_depth must be >= 1.")
        if self.min_samples_leaf < 1:
            raise ValueError("min_samples_leaf must be >= 1.")
        if self.n_estimators < 1:
            raise ValueError("n_estimators must be >= 1.")
        if self.permutation_repeats < 1:
            raise ValueError("permutation_repeats must be >= 1.")
        if self.vif_threshold <= 0:
            raise ValueError("vif_threshold must be > 0.")
        if self.max_interaction_features < 2:
            raise ValueError("max_interaction_features must be >= 2.")
        if self.max_three_way_combinations < 1:
            raise ValueError("max_three_way_combinations must be >= 1.")
        if self.n_clusters < 2:
            raise ValueError("n_clusters must be >= 2.")
        if self.max_clustering_features < 2:
            raise ValueError("max_clustering_features must be >= 2.")
        if not 0 < self.contamination <= 0.5:
            raise ValueError("contamination must be in the range (0, 0.5].")
        if self.summary_top_n < 1:
            raise ValueError("summary_top_n must be >= 1.")
        if self.min_group_size < 1:
            raise ValueError("min_group_size must be >= 1.")

    def _initialize_features(self) -> None:
        if self.df is None:
            raise ValueError("DataFrame must be loaded before feature initialization.")
        self.classified_columns = self.load_classified_columns()
        self.numeric_features = self._get_numeric_features()
        self.categorical_features = self._get_categorical_features()
        self.logger.info(
            "Initialized multivariate features: %d numeric, %d categorical.",
            len(self.numeric_features), len(self.categorical_features),
        )

    # ------------------------------------------------------------------
    # Classified columns
    # ------------------------------------------------------------------

    def load_classified_columns(self) -> dict[str, Any]:
        try:
            if not self.CLASSIFIED_COLUMNS_PATH.exists():
                self.logger.warning(
                    "classified_columns.json not found. Falling back to DataFrame dtypes."
                )
                return {}
            with self.CLASSIFIED_COLUMNS_PATH.open("r", encoding="utf-8") as file:
                data = json.load(file)
            if not isinstance(data, dict):
                self.logger.warning("classified_columns.json does not contain a dictionary.")
                return {}
            return data
        except Exception as exc:
            self.logger.warning("Could not load classified columns: %s", exc)
            return {}

    @staticmethod
    def _safe_float(value: Any) -> float | None:
        try:
            if value is None or pd.isna(value):
                return None
            return float(value)
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _is_id_like(df: pd.DataFrame, column: str) -> bool:
        normalized = column.strip().lower()
        if normalized in MultivariateEDA.ID_NAME_TOKENS:
            return True
        if normalized.endswith("_id"):
            return True
        if normalized.endswith("id") and len(normalized) > 2:
            return True
        try:
            unique_ratio = df[column].nunique(dropna=False) / max(len(df), 1)
            if unique_ratio >= 0.98 and df[column].nunique() > 20:
                return True
        except Exception:
            return False
        return False

    def _extract_classified(self, category: str) -> list[str]:
        if not self.classified_columns:
            return []
        value = self.classified_columns.get(category, [])
        if isinstance(value, dict):
            value = value.get("columns", [])
        if not isinstance(value, list):
            return []
        return [col for col in value if isinstance(col, str)]

    def _get_numeric_features(self) -> list[str]:
        if self.df is None:
            return []
        classified = self._extract_classified("numeric")
        if classified:
            numeric = [
                col for col in classified
                if col in self.df.columns and pd.api.types.is_numeric_dtype(self.df[col])
            ]
        else:
            numeric = [
                col for col in self.df.columns
                if pd.api.types.is_numeric_dtype(self.df[col])
            ]
        numeric = [col for col in numeric if not self._is_id_like(self.df, col)]
        return list(dict.fromkeys(numeric))

    def _get_categorical_features(self) -> list[str]:
        if self.df is None:
            return []
        classified = self._extract_classified("categorical")
        if classified:
            categorical = [col for col in classified if col in self.df.columns]
        else:
            categorical = [
                col for col in self.df.columns
                if (
                    pd.api.types.is_object_dtype(self.df[col])
                    or pd.api.types.is_categorical_dtype(self.df[col])
                    or pd.api.types.is_bool_dtype(self.df[col])
                )
            ]
        categorical = [col for col in categorical if not self._is_id_like(self.df, col)]
        return list(dict.fromkeys(categorical))

    # ------------------------------------------------------------------
    # General helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _top_n(data: list[dict[str, Any]], n: int) -> list[dict[str, Any]]:
        return data[:n]

    @staticmethod
    def _empty_result(columns: list[str]) -> pd.DataFrame:
        return pd.DataFrame(columns=columns)

    def _prepare_numeric_data(self, columns: list[str]) -> pd.DataFrame:
        if self.df is None:
            return pd.DataFrame()
        valid_columns = [col for col in columns if col in self.df.columns]
        if not valid_columns:
            return pd.DataFrame()
        data = self.df[valid_columns].apply(pd.to_numeric, errors="coerce")
        return data.replace([np.inf, -np.inf], np.nan).dropna()

    # ------------------------------------------------------------------
    # Target
    # ------------------------------------------------------------------

    def _detect_problem_type(self, target: pd.Series) -> str:
        if (
            pd.api.types.is_object_dtype(target)
            or pd.api.types.is_categorical_dtype(target)
            or pd.api.types.is_bool_dtype(target)
        ):
            return "classification"
        if pd.api.types.is_integer_dtype(target) and target.nunique() <= 20:
            return "classification"
        if target.nunique() <= 2:
            return "classification"
        return "regression"

    def _prepare_target(self) -> tuple[pd.Series, str] | None:
        if self.df is None or not self.target:
            return None
        if self.target not in self.df.columns:
            self.logger.warning("Target column not found: %s", self.target)
            return None
        target = self.df[self.target].copy()
        target = target.loc[target.notna()]
        if target.empty:
            return None
        return target, self._detect_problem_type(target)

    def _split_target(self, X: pd.DataFrame, y: pd.Series) -> tuple[pd.DataFrame, pd.Series]:
        combined = X.copy()
        combined["_target_"] = y
        combined = combined.dropna(subset=["_target_"])
        y_clean = combined.pop("_target_")
        return combined, y_clean

    def _classification_target_rate(self, y: pd.Series) -> float | None:
        if y.empty:
            return None
        encoded, _ = pd.factorize(y)
        if len(np.unique(encoded)) != 2:
            return None
        return float(np.mean(encoded))

    # ------------------------------------------------------------------
    # VIF
    # ------------------------------------------------------------------

    @track_performance
    def generate_vif(self) -> pd.DataFrame:
        columns = self.numeric_features.copy()
        if self.target in columns:
            columns.remove(self.target)
        empty_cols = ["feature", "r_squared", "vif", "high_vif"]
        if len(columns) < 2:
            return self._empty_result(empty_cols)

        data = self._prepare_numeric_data(columns)
        if len(data) < 3:
            return self._empty_result(empty_cols)

        results: list[dict[str, Any]] = []
        for feature in columns:
            others = [col for col in columns if col != feature]
            if not others:
                continue
            subset = data[[feature] + others].dropna()
            if len(subset) < 3:
                continue
            X, y = subset[others], subset[feature]
            try:
                model = LinearRegression()
                model.fit(X, y)
                r_squared = float(model.score(X, y))
                vif = np.inf if r_squared >= 0.999999 else 1.0 / (1.0 - r_squared)
                results.append({
                    "feature": feature,
                    "r_squared": self._safe_float(r_squared),
                    "vif": None if np.isinf(vif) else self._safe_float(vif),
                    "high_vif": bool(np.isinf(vif) or vif >= self.vif_threshold),
                })
            except Exception as exc:
                self.logger.warning("VIF failed for %s: %s", feature, exc)

        return pd.DataFrame(results)

    # ------------------------------------------------------------------
    # PCA
    # ------------------------------------------------------------------

    @track_performance
    def generate_pca(self) -> dict[str, pd.DataFrame]:
        columns = self.numeric_features.copy()
        if self.target in columns:
            columns.remove(self.target)
        data = self._prepare_numeric_data(columns)
        empty = {"pca_variance": pd.DataFrame(), "pca_loadings": pd.DataFrame()}

        if data.shape[0] < 2 or data.shape[1] < 2:
            return empty

        n_components = min(data.shape[0], data.shape[1])
        if self.max_pca_components is not None:
            n_components = min(n_components, self.max_pca_components)
        if n_components < 1:
            return empty

        try:
            scaled = StandardScaler().fit_transform(data)
            pca = PCA(n_components=n_components, random_state=self.random_state)
            pca.fit_transform(scaled)

            explained = pca.explained_variance_ratio_
            cumulative = np.cumsum(explained)
            variance_rows = [
                {
                    "component": f"PC{i + 1}",
                    "explained_variance_ratio": float(v),
                    "explained_variance_percentage": float(v * 100),
                    "cumulative_variance_ratio": float(cumulative[i]),
                    "cumulative_variance_percentage": float(cumulative[i] * 100),
                }
                for i, v in enumerate(explained)
            ]

            loading_rows = [
                {
                    "component": f"PC{ci + 1}",
                    "feature": feature,
                    "loading": float(loading),
                    "absolute_loading": float(abs(loading)),
                }
                for ci in range(pca.components_.shape[0])
                for feature, loading in zip(data.columns, pca.components_[ci])
            ]

            return {
                "pca_variance": pd.DataFrame(variance_rows),
                "pca_loadings": pd.DataFrame(loading_rows),
            }
        except Exception as exc:
            self.logger.warning("PCA failed: %s", exc)
            return empty

    # ------------------------------------------------------------------
    # Decision tree
    # ------------------------------------------------------------------

    def _prepare_model_features(self) -> tuple[pd.DataFrame, pd.Series, str] | None:
        if self.df is None or not self.target or self.target not in self.df.columns:
            return None

        numeric = [col for col in self.numeric_features if col != self.target]
        categorical = [
            col for col in self.categorical_features
            if col != self.target and self.df[col].nunique(dropna=True) <= 20
        ]
        features = numeric + categorical
        if not features:
            return None

        X = self.df[features].copy()
        y = self.df[self.target].copy()
        valid = y.notna()
        X, y = X.loc[valid], y.loc[valid]
        if len(X) < 10:
            return None

        numeric_present = [col for col in numeric if col in X.columns]
        categorical_present = [col for col in categorical if col in X.columns]
        transformers = []
        if numeric_present:
            transformers.append(("numeric", StandardScaler(), numeric_present))
        if categorical_present:
            transformers.append((
                "categorical",
                OneHotEncoder(handle_unknown="ignore", sparse_output=False),
                categorical_present,
            ))
        if not transformers:
            return None

        try:
            preprocessor = ColumnTransformer(transformers=transformers, remainder="drop")
            X_transformed = preprocessor.fit_transform(X)
            feature_names = preprocessor.get_feature_names_out()
            X_result = pd.DataFrame(X_transformed, columns=feature_names, index=X.index)
            return X_result, y, self._detect_problem_type(y)
        except Exception as exc:
            self.logger.warning("Model feature preparation failed: %s", exc)
            return None

    def _extract_tree_decision_points(
        self, model: Any, feature_names: list[str]
    ) -> list[dict[str, Any]]:
        tree = model.tree_
        results = []
        for node_id in range(tree.node_count):
            feature_index = tree.feature[node_id]
            if feature_index < 0 or feature_index >= len(feature_names):
                continue
            results.append({
                "node": node_id,
                "depth": self._tree_node_depth(tree, node_id),
                "feature": feature_names[feature_index],
                "threshold": self._safe_float(tree.threshold[node_id]),
                "samples": int(tree.n_node_samples[node_id]),
            })
        return results

    @staticmethod
    def _tree_node_depth(tree: Any, node_id: int) -> int:
        parents: dict[int, int] = {}
        for parent in range(tree.node_count):
            left, right = tree.children_left[parent], tree.children_right[parent]
            if left != -1:
                parents[left] = parent
            if right != -1:
                parents[right] = parent
        depth, current = 0, node_id
        while current in parents:
            depth += 1
            current = parents[current]
        return depth

    def _extract_tree_paths(
        self, model: Any, feature_names: list[str]
    ) -> list[dict[str, Any]]:
        tree = model.tree_
        paths: list[dict[str, Any]] = []

        def traverse(node_id: int, conditions: list[str]) -> None:
            left, right = tree.children_left[node_id], tree.children_right[node_id]
            if left == -1 and right == -1:
                paths.append({
                    "leaf": node_id,
                    "conditions": " AND ".join(conditions) if conditions else "ROOT",
                    "samples": int(tree.n_node_samples[node_id]),
                })
                return
            feature_index = tree.feature[node_id]
            if feature_index < 0:
                return
            feature = feature_names[feature_index]
            threshold = tree.threshold[node_id]
            traverse(left, conditions + [f"{feature} <= {threshold:.6f}"])
            traverse(right, conditions + [f"{feature} > {threshold:.6f}"])

        try:
            traverse(0, [])
        except Exception as exc:
            self.logger.warning("Tree path extraction failed: %s", exc)
        return paths

    def _extract_leaf_profiles(self, model: Any) -> list[dict[str, Any]]:
        tree = model.tree_
        profiles = []
        for node_id in range(tree.node_count):
            if tree.children_left[node_id] != -1 or tree.children_right[node_id] != -1:
                continue
            profiles.append({
                "leaf": node_id,
                "samples": int(tree.n_node_samples[node_id]),
                "value": tree.value[node_id].tolist(),
            })
        return profiles

    @track_performance
    def generate_decision_tree(self) -> dict[str, pd.DataFrame]:
        empty = {
            "tree_metrics": pd.DataFrame(),
            "tree_decision_points": pd.DataFrame(),
            "tree_decision_paths": pd.DataFrame(),
            "tree_leaf_profiles": pd.DataFrame(),
            "tree_feature_importance": pd.DataFrame(),
        }
        prepared = self._prepare_model_features()
        if prepared is None:
            return empty
        X, y, problem_type = prepared
        if y.nunique() < 2:
            return empty

        try:
            X_train, X_test, y_train, y_test = train_test_split(
                X, y, test_size=self.test_size, random_state=self.random_state
            )
            if problem_type == "classification":
                model = DecisionTreeClassifier(
                    max_depth=self.max_depth, min_samples_leaf=self.min_samples_leaf,
                    random_state=self.random_state,
                )
            else:
                model = DecisionTreeRegressor(
                    max_depth=self.max_depth, min_samples_leaf=self.min_samples_leaf,
                    random_state=self.random_state,
                )
            model.fit(X_train, y_train)
            predictions = model.predict(X_test)

            if problem_type == "classification":
                metric_rows = [{
                    "problem_type": problem_type,
                    "metric": "accuracy",
                    "value": float(accuracy_score(y_test, predictions)),
                }]
            else:
                metric_rows = [
                    {"problem_type": problem_type, "metric": "r2",
                     "value": float(r2_score(y_test, predictions))},
                    {"problem_type": problem_type, "metric": "mae",
                     "value": float(mean_absolute_error(y_test, predictions))},
                    {"problem_type": problem_type, "metric": "rmse",
                     "value": float(np.sqrt(mean_squared_error(y_test, predictions)))},
                ]

            feature_names = list(X.columns)
            decision_points = self._extract_tree_decision_points(model, feature_names)
            paths = self._extract_tree_paths(model, feature_names)
            leaf_profiles = self._extract_leaf_profiles(model)

            importance_rows = [
                {"feature": f, "importance": float(imp)}
                for f, imp in zip(feature_names, model.feature_importances_)
            ]
            importance_rows.sort(key=lambda r: r["importance"], reverse=True)

            return {
                "tree_metrics": pd.DataFrame(metric_rows),
                "tree_decision_points": pd.DataFrame(
                    self._top_n(decision_points, self.summary_top_n)
                ),
                "tree_decision_paths": pd.DataFrame(self._top_n(paths, self.summary_top_n)),
                "tree_leaf_profiles": pd.DataFrame(
                    self._top_n(leaf_profiles, self.summary_top_n)
                ),
                "tree_feature_importance": pd.DataFrame(
                    self._top_n(importance_rows, self.summary_top_n)
                ),
            }
        except Exception as exc:
            self.logger.warning("Decision tree analysis failed: %s", exc)
            return empty

    # ------------------------------------------------------------------
    # Linear / logistic models
    # ------------------------------------------------------------------

    @track_performance
    def generate_models(self) -> dict[str, pd.DataFrame]:
        empty = {
            "model_metrics": pd.DataFrame(),
            "model_coefficients": pd.DataFrame(),
            "model_residuals": pd.DataFrame(),
        }
        prepared = self._prepare_model_features()
        if prepared is None:
            return empty
        X, y, problem_type = prepared
        if y.nunique() < 2:
            return empty

        try:
            X_train, X_test, y_train, y_test = train_test_split(
                X, y, test_size=self.test_size, random_state=self.random_state
            )

            if problem_type == "classification":
                model = LogisticRegression(max_iter=1000, random_state=self.random_state)
                model.fit(X_train, y_train)
                predictions = model.predict(X_test)
                metrics = [{
                    "problem_type": problem_type,
                    "metric": "accuracy",
                    "value": float(accuracy_score(y_test, predictions)),
                }]
                coefficient_rows = [
                    {
                        "class": class_index,
                        "feature": feature,
                        "coefficient": float(coef),
                        "odds_ratio": float(np.exp(coef)),
                    }
                    for class_index, class_coefs in enumerate(model.coef_)
                    for feature, coef in zip(X.columns, class_coefs)
                ]
                residuals = pd.DataFrame()
            else:
                model = LinearRegression()
                model.fit(X_train, y_train)
                predictions = model.predict(X_test)
                metrics = [
                    {"problem_type": problem_type, "metric": "r2",
                     "value": float(r2_score(y_test, predictions))},
                    {"problem_type": problem_type, "metric": "mae",
                     "value": float(mean_absolute_error(y_test, predictions))},
                    {"problem_type": problem_type, "metric": "rmse",
                     "value": float(np.sqrt(mean_squared_error(y_test, predictions)))},
                ]
                coefficient_rows = [
                    {"feature": f, "coefficient": float(c)}
                    for f, c in zip(X.columns, model.coef_)
                ]
                residuals = pd.DataFrame({
                    "actual": y_test.to_numpy(),
                    "predicted": predictions,
                    "residual": y_test.to_numpy() - predictions,
                })

            coefficient_rows.sort(key=lambda r: abs(r.get("coefficient", 0)), reverse=True)
            return {
                "model_metrics": pd.DataFrame(metrics),
                "model_coefficients": pd.DataFrame(
                    self._top_n(coefficient_rows, self.summary_top_n)
                ),
                "model_residuals": residuals.head(self.summary_top_n),
            }
        except Exception as exc:
            self.logger.warning("Regression/classification model failed: %s", exc)
            return empty

    # ------------------------------------------------------------------
    # Random Forest
    # ------------------------------------------------------------------

    @track_performance
    def generate_random_forest(self) -> dict[str, pd.DataFrame]:
        empty = {
            "rf_metrics": pd.DataFrame(),
            "rf_feature_importance": pd.DataFrame(),
            "rf_permutation_importance": pd.DataFrame(),
        }
        prepared = self._prepare_model_features()
        if prepared is None:
            return empty
        X, y, problem_type = prepared
        if y.nunique() < 2:
            return empty

        try:
            X_train, X_test, y_train, y_test = train_test_split(
                X, y, test_size=self.test_size, random_state=self.random_state
            )
            if problem_type == "classification":
                model = RandomForestClassifier(
                    n_estimators=self.n_estimators, random_state=self.random_state,
                    n_jobs=-1, oob_score=True,
                )
            else:
                model = RandomForestRegressor(
                    n_estimators=self.n_estimators, random_state=self.random_state,
                    n_jobs=-1, oob_score=True,
                )
            model.fit(X_train, y_train)
            predictions = model.predict(X_test)

            if problem_type == "classification":
                metrics = [
                    {"problem_type": problem_type, "metric": "accuracy",
                     "value": float(accuracy_score(y_test, predictions))},
                    {"problem_type": problem_type, "metric": "oob_score",
                     "value": self._safe_float(model.oob_score_)},
                ]
            else:
                metrics = [
                    {"problem_type": problem_type, "metric": "r2",
                     "value": float(r2_score(y_test, predictions))},
                    {"problem_type": problem_type, "metric": "mae",
                     "value": float(mean_absolute_error(y_test, predictions))},
                    {"problem_type": problem_type, "metric": "rmse",
                     "value": float(np.sqrt(mean_squared_error(y_test, predictions)))},
                    {"problem_type": problem_type, "metric": "oob_score",
                     "value": self._safe_float(model.oob_score_)},
                ]

            importance_rows = [
                {"feature": f, "importance": float(imp)}
                for f, imp in zip(X.columns, model.feature_importances_)
            ]
            importance_rows.sort(key=lambda r: r["importance"], reverse=True)

            try:
                perm = permutation_importance(
                    model, X_test, y_test, n_repeats=self.permutation_repeats,
                    random_state=self.random_state, n_jobs=-1,
                )
                permutation_rows = [
                    {
                        "feature": f,
                        "importance_mean": float(mean),
                        "importance_std": float(std),
                    }
                    for f, mean, std in zip(
                        X.columns, perm.importances_mean, perm.importances_std
                    )
                ]
                permutation_rows.sort(key=lambda r: r["importance_mean"], reverse=True)
            except Exception as exc:
                self.logger.warning("Permutation importance failed: %s", exc)
                permutation_rows = []

            return {
                "rf_metrics": pd.DataFrame(metrics),
                "rf_feature_importance": pd.DataFrame(
                    self._top_n(importance_rows, self.summary_top_n)
                ),
                "rf_permutation_importance": pd.DataFrame(
                    self._top_n(permutation_rows, self.summary_top_n)
                ),
            }
        except Exception as exc:
            self.logger.warning("Random Forest analysis failed: %s", exc)
            return empty

    # ------------------------------------------------------------------
    # Two-way / three-way interactions
    # ------------------------------------------------------------------

    @track_performance
    def generate_two_way_interactions(self) -> pd.DataFrame:
        if self.df is None or not self.target or self.target not in self.df.columns:
            return pd.DataFrame()

        features = [
            f for f in self.numeric_features if f != self.target
        ][: self.max_interaction_features]
        if len(features) < 2:
            return pd.DataFrame()

        target_series = self.df[self.target]
        if not pd.api.types.is_numeric_dtype(target_series):
            target_numeric = pd.factorize(target_series)[0].astype(float)
        else:
            target_numeric = pd.to_numeric(target_series, errors="coerce")

        rows: list[dict[str, Any]] = []
        for feature_a, feature_b in itertools.combinations(features, 2):
            try:
                subset = pd.DataFrame({
                    "feature_a": pd.to_numeric(self.df[feature_a], errors="coerce"),
                    "feature_b": pd.to_numeric(self.df[feature_b], errors="coerce"),
                    "target": target_numeric,
                }).dropna()
                if len(subset) < self.min_group_size:
                    continue

                subset["a_bin"] = pd.qcut(subset["feature_a"], q=3, duplicates="drop")
                subset["b_bin"] = pd.qcut(subset["feature_b"], q=3, duplicates="drop")
                grouped = (
                    subset.groupby(["a_bin", "b_bin"], observed=True)["target"]
                    .agg(["mean", "count"]).reset_index()
                )
                grouped = grouped[grouped["count"] >= self.min_group_size]
                if grouped.empty:
                    continue

                target_difference = float(grouped["mean"].max()) - float(grouped["mean"].min())
                max_group = grouped.loc[grouped["count"].idxmax()]
                rows.append({
                    "feature_a": feature_a,
                    "feature_b": feature_b,
                    "groups_evaluated": len(grouped),
                    "target_min": float(grouped["mean"].min()),
                    "target_max": float(grouped["mean"].max()),
                    "target_difference": target_difference,
                    "largest_group": str(max_group["count"]),
                })
            except Exception as exc:
                self.logger.warning(
                    "Two-way interaction failed for %s, %s: %s", feature_a, feature_b, exc
                )

        rows.sort(key=lambda r: r["target_difference"], reverse=True)
        return pd.DataFrame(self._top_n(rows, self.summary_top_n))

    @track_performance
    def generate_three_way_interactions(self) -> pd.DataFrame:
        if self.df is None or not self.target or self.target not in self.df.columns:
            return pd.DataFrame()

        features = [f for f in self.numeric_features if f != self.target]
        if len(features) < 3:
            return pd.DataFrame()
        features = features[: self.max_interaction_features]

        rows: list[dict[str, Any]] = []
        for feature_a, feature_b, feature_c in itertools.islice(
            itertools.combinations(features, 3), self.max_three_way_combinations
        ):
            try:
                target_series = self.df[self.target]
                if pd.api.types.is_numeric_dtype(target_series):
                    target = pd.to_numeric(target_series, errors="coerce")
                else:
                    target = pd.Series(
                        pd.factorize(target_series)[0], index=target_series.index, dtype=float
                    )

                subset = pd.DataFrame({
                    "a": pd.to_numeric(self.df[feature_a], errors="coerce"),
                    "b": pd.to_numeric(self.df[feature_b], errors="coerce"),
                    "c": pd.to_numeric(self.df[feature_c], errors="coerce"),
                    "target": target,
                }).dropna()
                if len(subset) < self.min_group_size:
                    continue

                subset["a_bin"] = pd.qcut(subset["a"], q=2, duplicates="drop")
                subset["b_bin"] = pd.qcut(subset["b"], q=2, duplicates="drop")
                subset["c_bin"] = pd.qcut(subset["c"], q=2, duplicates="drop")
                grouped = (
                    subset.groupby(["a_bin", "b_bin", "c_bin"], observed=True)["target"]
                    .agg(["mean", "count"]).reset_index()
                )
                grouped = grouped[grouped["count"] >= self.min_group_size]
                if grouped.empty:
                    continue

                target_min = float(grouped["mean"].min())
                target_max = float(grouped["mean"].max())
                rows.append({
                    "feature_a": feature_a,
                    "feature_b": feature_b,
                    "feature_c": feature_c,
                    "groups_evaluated": len(grouped),
                    "target_min": target_min,
                    "target_max": target_max,
                    "target_difference": target_max - target_min,
                    "max_group_count": int(grouped["count"].max()),
                })
            except Exception as exc:
                self.logger.warning(
                    "Three-way interaction failed for %s, %s, %s: %s",
                    feature_a, feature_b, feature_c, exc,
                )

        rows.sort(key=lambda r: r["target_difference"], reverse=True)
        return pd.DataFrame(self._top_n(rows, self.summary_top_n))

    # ------------------------------------------------------------------
    # Clustering
    # ------------------------------------------------------------------

    @track_performance
    def generate_clustering(self) -> dict[str, pd.DataFrame]:
        empty = {
            "cluster_metrics": pd.DataFrame(),
            "cluster_profile": pd.DataFrame(),
            "cluster_assignments": pd.DataFrame(),
        }
        columns = [c for c in self.numeric_features if c != self.target][
            : self.max_clustering_features
        ]
        data = self._prepare_numeric_data(columns)
        if len(data) < 3 or data.shape[1] < 2:
            return empty

        try:
            scaled = StandardScaler().fit_transform(data)
            max_k = min(8, self.n_clusters, len(data) - 1)
            if max_k < 2:
                return empty

            metric_rows = []
            models: dict[int, KMeans] = {}
            for k in range(2, max_k + 1):
                model = KMeans(n_clusters=k, random_state=self.random_state, n_init=10)
                labels = model.fit_predict(scaled)
                if len(np.unique(labels)) < 2:
                    continue
                models[k] = model
                metric_rows.append({
                    "k": k,
                    "silhouette_score": float(silhouette_score(scaled, labels)),
                    "calinski_harabasz_score": float(calinski_harabasz_score(scaled, labels)),
                    "davies_bouldin_score": float(davies_bouldin_score(scaled, labels)),
                })

            if not metric_rows:
                return empty

            best = max(metric_rows, key=lambda r: r["silhouette_score"])
            best_k = best["k"]
            best_model = models[best_k]
            labels = best_model.labels_

            profile_data = data.copy()
            profile_data["cluster"] = labels
            profile = profile_data.groupby("cluster").mean(numeric_only=True).reset_index()
            assignments = pd.DataFrame({"row_index": data.index, "cluster": labels})
            metrics_df = pd.DataFrame(metric_rows)
            metrics_df["best_k"] = metrics_df["k"] == best_k

            return {
                "cluster_metrics": metrics_df,
                "cluster_profile": profile,
                "cluster_assignments": assignments.head(self.summary_top_n),
            }
        except Exception as exc:
            self.logger.warning("Clustering analysis failed: %s", exc)
            return empty

    # ------------------------------------------------------------------
    # Outliers
    # ------------------------------------------------------------------

    @track_performance
    def generate_outliers(self) -> dict[str, pd.DataFrame]:
        columns = self.numeric_features.copy()
        if self.target in columns:
            columns.remove(self.target)
        data = self._prepare_numeric_data(columns)
        empty = {
            "isolation_forest_outliers": pd.DataFrame(),
            "mahalanobis_outliers": pd.DataFrame(),
        }
        if len(data) < 10 or data.shape[1] < 2:
            return empty

        isolation_rows: list[dict[str, Any]] = []
        mahalanobis_rows: list[dict[str, Any]] = []

        try:
            scaled = StandardScaler().fit_transform(data)
            model = IsolationForest(
                contamination=self.contamination, random_state=self.random_state
            )
            predictions = model.fit_predict(scaled)
            scores = model.decision_function(scaled)
            outlier_indices = data.index[predictions == -1]
            for index in outlier_indices:
                position = data.index.get_loc(index)
                isolation_rows.append({
                    "row_index": index,
                    "outlier": True,
                    "anomaly_score": float(scores[position]),
                })
            isolation_rows.sort(key=lambda r: r["anomaly_score"])
        except Exception as exc:
            self.logger.warning("Isolation Forest failed: %s", exc)

        try:
            values = data.to_numpy(dtype=float)
            mean = np.mean(values, axis=0)
            covariance = np.cov(values, rowvar=False)
            inverse_covariance = np.linalg.pinv(covariance)
            centered = values - mean
            distances = np.einsum("ij,jk,ik->i", centered, inverse_covariance, centered)
            threshold = chi2.ppf(self.MAHALANOBIS_CONFIDENCE, df=data.shape[1])
            mask = distances > threshold
            for index, distance in zip(data.index[mask], distances[mask]):
                mahalanobis_rows.append({
                    "row_index": index,
                    "mahalanobis_distance": float(distance),
                    "chi_square_threshold": float(threshold),
                    "outlier": True,
                })
            mahalanobis_rows.sort(key=lambda r: r["mahalanobis_distance"], reverse=True)
        except Exception as exc:
            self.logger.warning("Mahalanobis outlier analysis failed: %s", exc)

        return {
            "isolation_forest_outliers": pd.DataFrame(
                self._top_n(isolation_rows, self.summary_top_n)
            ),
            "mahalanobis_outliers": pd.DataFrame(
                self._top_n(mahalanobis_rows, self.summary_top_n)
            ),
        }

    # ------------------------------------------------------------------
    # Excel output
    # ------------------------------------------------------------------

    @staticmethod
    def _serialize_value(value: Any) -> Any:
        if isinstance(value, (dict, list, tuple)):
            try:
                return json.dumps(value, default=str)
            except Exception:
                return str(value)
        if isinstance(value, (np.integer, np.floating, np.bool_)):
            return value.item()
        if pd.isna(value):
            return None
        return value

    def _prepare_dataframe_for_excel(self, df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return df.copy()
        result = df.copy()
        for column in result.columns:
            result[column] = result[column].map(self._serialize_value)
        return result

    def _write_dataframe_sheet(
        self, workbook: Workbook, sheet_name: str, df: pd.DataFrame
    ) -> None:
        safe_name = sheet_name[:31]
        if safe_name in workbook.sheetnames:
            del workbook[safe_name]
        worksheet = workbook.create_sheet(title=safe_name)
        df = self._prepare_dataframe_for_excel(df)

        if df.empty:
            worksheet.append(["No data available."])
            return

        worksheet.append(list(df.columns))
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for row in df.itertuples(index=False, name=None):
            worksheet.append(list(row))

        for column_cells in worksheet.columns:
            max_length = 0
            for cell in column_cells:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    continue
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                max(max_length + 2, 10), 50
            )

    @track_performance
    def save_workbook(self, results: dict[str, Any]) -> Path:
        try:
            self.output_dir.mkdir(parents=True, exist_ok=True)
            output_path = self.output_dir / self.WORKBOOK_NAME
            workbook = Workbook()
            default_sheet = workbook.active
            workbook.remove(default_sheet)

            for sheet_name, result in results.items():
                if isinstance(result, pd.DataFrame):
                    self._write_dataframe_sheet(workbook, sheet_name, result)
                elif isinstance(result, dict):
                    for nested_name, nested_result in result.items():
                        if isinstance(nested_result, pd.DataFrame):
                            self._write_dataframe_sheet(workbook, nested_name, nested_result)

            workbook.save(output_path)
            self.logger.info("Multivariate EDA workbook saved: %s", output_path)
            return output_path
        except Exception as exc:
            self.logger.error("Failed to save Multivariate EDA workbook: %s", exc)
            raise CustomException(exc) from exc

    # ------------------------------------------------------------------
    # Main execution
    # ------------------------------------------------------------------

    @track_performance
    def run_eda(self) -> dict[str, pd.DataFrame]:
        if self.df is None:
            raise ValueError("DataFrame is not loaded. Use run_from_file() or provide df.")
        if self.df.empty:
            raise ValueError("Cannot run Multivariate EDA on an empty DataFrame.")

        self.logger.info("Starting Multivariate EDA.")
        results: dict[str, Any] = {}

        try:
            results["vif"] = self.generate_vif()
        except Exception as exc:
            self.logger.warning("VIF generation failed: %s", exc)

        try:
            results.update(self.generate_pca())
        except Exception as exc:
            self.logger.warning("PCA generation failed: %s", exc)

        if self.target:
            try:
                results.update(self.generate_decision_tree())
            except Exception as exc:
                self.logger.warning("Decision tree generation failed: %s", exc)
            try:
                results.update(self.generate_models())
            except Exception as exc:
                self.logger.warning("Model generation failed: %s", exc)
            try:
                results.update(self.generate_random_forest())
            except Exception as exc:
                self.logger.warning("Random Forest generation failed: %s", exc)
            try:
                results["two_way_interactions"] = self.generate_two_way_interactions()
            except Exception as exc:
                self.logger.warning("Two-way interaction generation failed: %s", exc)
            try:
                results["three_way_interactions"] = self.generate_three_way_interactions()
            except Exception as exc:
                self.logger.warning("Three-way interaction generation failed: %s", exc)
        else:
            self.logger.info("No target supplied. Skipping target-dependent analyses.")

        try:
            results.update(self.generate_clustering())
        except Exception as exc:
            self.logger.warning("Clustering generation failed: %s", exc)
        try:
            results.update(self.generate_outliers())
        except Exception as exc:
            self.logger.warning("Outlier generation failed: %s", exc)

        output_path = self.save_workbook(results)
        self.logger.info("Multivariate EDA completed successfully.")
        self.logger.info("Workbook: %s", output_path)
        return results

    @track_performance
    def run_from_file(self, file_path: Path | str) -> dict[str, pd.DataFrame]:
        file_path = Path(file_path)
        try:
            if not file_path.exists():
                raise FileNotFoundError(f"Dataset not found: {file_path}")
            if not file_path.is_file():
                raise FileNotFoundError(f"Dataset path is not a file: {file_path}")

            self.logger.info("Loading dataset: %s", file_path)
            self.df = pd.read_csv(file_path)
            if self.df.empty:
                raise ValueError(f"Dataset is empty: {file_path}")

            self.logger.info(
                "Dataset loaded successfully: %d rows x %d columns.",
                self.df.shape[0], self.df.shape[1],
            )
            self._initialize_features()
            return self.run_eda()
        except Exception as exc:
            self.logger.error("Multivariate EDA failed for %s: %s", file_path, exc)
            raise CustomException(exc) from exc


@track_performance
def multivariate_stats(
    file_path: Path | str | None = None, target: str="Returned" 
) -> dict[str, pd.DataFrame]:
    """Run Multivariate EDA using the standard project dataset path."""
    if file_path is None:
        file_path = BASE_DIR / "src/Data/cleaned_ecommerce_dataset.csv"
    return MultivariateEDA(target=target).run_from_file(file_path=file_path)


def run_multivariate_eda(
    df: pd.DataFrame, target: str="Returned" 
) -> dict[str, pd.DataFrame]:
    """Run Multivariate EDA directly from an existing DataFrame."""
    return MultivariateEDA(df=df, target=target).run_eda()


if __name__ == "__main__":
    try:
        multivariate_stats()
    except Exception as exc:
        print(f"Multivariate EDA failed: {exc}")